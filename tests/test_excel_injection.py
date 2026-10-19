"""Excel formula injection: user/model text must never become formulas."""

from __future__ import annotations

import unittest
from io import BytesIO

from openpyxl import load_workbook

from enrich import CleanRow, RunStats, SourceRow, STATUS_REVIEW, write_workbook

PAYLOADS = [
    "=1+1",
    '=HYPERLINK("http://x","y")',
    "+cmd|' /C calc'!A0",
    "@SUM(1,2)",
    "-2+3",
    "+84912345678",
]


def _poisoned_row(text: str, row_id: int) -> CleanRow:
    source = SourceRow(
        row_id=row_id,
        raw_product=text,
        raw_price=text,
        customer_name=text,
        raw_phone=text,
        raw_address=text,
        raw_feedback=text,
        price_vnd=None,
        phone=text if text.startswith("+849") else None,
        product_empty=False,
    )
    return CleanRow(
        source=source,
        clean_name_vi=text,
        clean_name_en=text,
        brand=text,
        category_vi="Khác",
        category_en="Other",
        address=text,
        summary_vi=text,
        summary_en=text,
        confidence=0.2,
        status=STATUS_REVIEW,
        reasons=[text],
        address_needs_review=True,
    )


class ExcelInjectionTests(unittest.TestCase):
    def test_payloads_are_stored_as_text_not_formulas(self) -> None:
        rows = [_poisoned_row(text, index) for index, text in enumerate(PAYLOADS, start=1)]
        stats = RunStats(rows=len(rows), ok=0, needs_review=len(rows), output_path="memory")
        buffer = BytesIO()
        write_workbook(rows, buffer, stats, "injection.xlsx")
        buffer.seek(0)
        workbook = load_workbook(buffer)

        seen: set[str] = set()
        formulas: list[str] = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        formulas.append(f"{sheet.title}!{cell.coordinate}={cell.value!r}")
                    if isinstance(cell.value, str) and cell.value in PAYLOADS:
                        seen.add(cell.value)
                        self.assertEqual(cell.data_type, "s")
                        self.assertTrue(cell.quotePrefix)

        self.assertEqual(formulas, [])
        self.assertEqual(seen, set(PAYLOADS))
        self.assertIn("+84912345678", seen)


if __name__ == "__main__":
    unittest.main()
